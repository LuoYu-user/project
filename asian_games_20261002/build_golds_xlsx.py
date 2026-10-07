# -*- coding: utf-8 -*-
"""第20届亚运会（爱知·名古屋 2026）2026-10-02 金牌汇总 工作簿生成器

流程：拉取数据（已完成，来源见"数据来源"页）→ 生成本 xlsx。
透视页所有汇总数字均为 Excel 公式（COUNTIF/SUM/COUNTA 等），非手写常数。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.environ.get("TEMP", "/tmp"), "asian_games_20261002_golds.xlsx")

# 2026-10-02（第13比赛日）当日决出的 51 枚金牌明细
# 官方当日分布：拳击11 帆船5 摔跤5 柔道5 跆拳道4 场地自行车4
# 射箭3 皮划艇激流回旋2 柔术3 壁球2 网球2 电子竞技1 足球1 虚拟跆拳道1 沙滩排球1 曲棍球1 = 51
# (序号, 大项, 分项, 冠军代表团, 冠军/队伍, 备注)
ROWS = [
    # 拳拳 11（当日全部收官）
    (1,  "拳击", "男子55公斤级",     "日本",         "山口来生",                 "3:2 胜菲律宾帕阿兰，日本52年来该级别首金"),
    (2,  "拳击", "男子60公斤级",     "乌兹别克斯坦", "阿卜杜马利克·卡洛科霍夫",   ""),
    (3,  "拳击", "男子70公斤级",     "哈萨克斯坦",   "托列汗·萨贝尔汗",           "5:0 胜约旦伊沙伊什"),
    (4,  "拳击", "男子80公斤级",     "印度",         "安库什·潘加勒",             "4:0 胜韩国金旻成"),
    (5,  "拳击", "男子90公斤级",     "乌兹别克斯坦", "图拉别克·哈比布拉耶夫",      "5:0 胜中国韩雪贞"),
    (6,  "拳击", "男子90公斤以上级", "哈萨克斯坦",   "艾别克·奥拉尔巴耶",         ""),
    (7,  "拳击", "女子51公斤级",     "中国",         "吴愉",                     "5:0 完胜越南选手"),
    (8,  "拳击", "女子54公斤级",     "泰国",         "纳特尼查·琼蓬格朗",         "决赛胜朝鲜朴春美"),
    (9,  "拳击", "女子60公斤级",     "中国台北",     "林郁婷",                   "决赛胜中国杨成宇"),
    (10, "拳击", "女子65公斤级",     "印度",         "帕维恩·胡达",               "3:2 胜中国台北陈念琴"),
    (11, "拳击", "女子75公斤级",     "印度",         "洛夫利娜·博尔戈哈因",       "5:0 胜中国鲍子怡"),
    # 帆船 5
    (12, "帆船", "男子ILCA 7级",     "韩国",         "河智旻",                   ""),
    (13, "帆船", "女子49erFX级",     "日本",         "田中美咲/长永濑世拉",       ""),
    (14, "帆船", "混合470级",        "日本",         "磯崎哲也/关友里惠",         ""),
    (15, "帆船", "男子iQFoil级",     "中国",         "毕坤",                     "10月3日风小取消，按2日末轮名次判金"),
    (16, "帆船", "女子iQFoil级",     "中国",         "李文琪",                   "10月3日风小取消，按2日末轮名次判金"),
    # 摔跤-自由式 5
    (17, "摔跤-自由式", "男子65公斤级",  "印度", "苏吉特·卡尔卡尔",     "10:9 逆转胜塔吉克斯坦库迪耶夫"),
    (18, "摔跤-自由式", "男子86公斤级",  "伊朗", "穆罕默德·诺霍迪",     "9:6 胜日本石黑隼人"),
    (19, "摔跤-自由式", "男子125公斤级", "伊朗", "阿米尔侯赛因·扎雷",   "12:2 胜中国布和额尔顿，卫冕"),
    (20, "摔跤-自由式", "女子50公斤级",  "日本", "须崎优衣",             "决赛零封朝鲜选手"),
    (21, "摔跤-自由式", "女子76公斤级",  "日本", "镜和翔",               "5:1 胜哈萨克斯坦选手"),
    # 柔道 5
    (22, "柔道", "男子90公斤级",     "日本",         "冈田陆",                   "金分技取胜"),
    (23, "柔道", "男子100公斤级",    "日本",         "增地辽太郎",               "一本取胜"),
    (24, "柔道", "女子78公斤级",     "日本",         "杉村美寿希",               "反则胜韩国金敏周"),
    (25, "柔道", "女子78公斤以上级", "中国",         "牛馨然",                   ""),
    (26, "柔道", "男子100公斤以上级", "乌兹别克斯坦", "图罗博耶夫·穆扎法尔别克",   ""),
    # 跆拳道 4
    (27, "跆拳道", "男子-58公斤级", "乌兹别克斯坦", "贾苏尔贝克·贾伊苏诺夫",     ""),
    (28, "跆拳道", "男子-80公斤级", "哈萨克斯坦",   "塔米尔兰·特勒乌列斯",       ""),
    (29, "跆拳道", "女子-49公斤级", "中国",         "马静月",                   ""),
    (30, "跆拳道", "女子-67公斤级", "伊朗",         "萨加尔·莫拉迪",             ""),
    # 场地自行车 4
    (31, "场地自行车", "男子凯林赛",   "马来西亚", "穆赫德·阿兹祖哈斯尼·阿旺", "决赛夺冠"),
    (32, "场地自行车", "男子麦迪逊赛", "日本",     "窪木一茂/儿岛直树",         "76分夺金"),
    (33, "场地自行车", "女子争先赛",   "中国",     "苑丽颖",                   "胜日本佐藤水菜"),
    (34, "场地自行车", "女子全能赛",   "日本",     "梶原悠未",                 "136分，三连冠"),
    # 射箭 3
    (35, "射箭", "反曲弓混合团体", "中国", "李梦麒/朱静怡", "5:3 胜印度，获洛杉矶奥运席位"),
    (36, "射箭", "反曲弓男子团体", "韩国", "韩国男子团体",   "5:1 胜印度"),
    (37, "射箭", "反曲弓女子团体", "印度", "安基塔·巴卡特/蒂尔蒂·夏尔玛/库姆库姆·莫霍德", "5:3 胜韩国，印度首夺该项目亚运金"),
    # 皮划艇激流回旋 2
    (38, "皮划艇激流回旋", "男子皮艇越野", "日本", "田中雄己", ""),
    (39, "皮划艇激流回旋", "女子皮艇越野", "中国", "陈琳",   "1分28秒72 刷赛会纪录"),
    # 柔术 3
    (40, "柔术", "男子-77公斤级", "阿联酋",     "马赫迪·阿劳拉齐",   "胜约旦贾布里"),
    (41, "柔术", "女子-52公斤级", "阿联酋",     "阿斯玛·胡斯尼",     "胜韩国任彦珠，成功卫冕"),
    (42, "柔术", "男子-94公斤级", "沙特阿拉伯", "奥马尔·塔里克·纳达", "决赛胜韩国朴成厚"),
    # 壁球 2
    (43, "壁球", "男子团体", "马来西亚", "马来西亚男子团体", "决赛胜中国香港"),
    (44, "壁球", "女子团体", "马来西亚", "马来西亚女子团体", "决赛胜中国香港，马来西亚包办双金"),
    # 网球 2
    (45, "网球", "男子双打",   "日本", "楠原悠介/中川舜祐", "6-3 7-6 胜韩国组合"),
    (46, "网球", "女子单打",   "中国", "王曦雨",           "决赛夺金"),
    # 电子竞技 1
    (47, "电子竞技", "英雄联盟", "韩国", "韩国队", "决赛胜中国台北"),
    # 足球 1
    (48, "足球", "女子足球", "日本", "日本女足", "2-2（点球5-3）胜朝鲜，亚运三连冠"),
    # 虚拟跆拳道 1
    (49, "跆拳道-虚拟", "混合个人", "菲律宾", "杰乌斯·加布里埃尔·亚佩", "2:0 胜越南选手"),
    # 沙滩排球 1
    (50, "沙滩排球", "女子", "泰国", "塔拉瓦迪·纳拉蓬/沃拉皮拉查亚功·功普沙塔翁", "决赛夺金"),
    # 曲棍球 1
    (51, "曲棍球", "女子曲棍球", "印度", "印度女曲", "1:0 胜中国，时隔44年再夺亚运金，获洛杉矶奥运席位"),
]

SPORT_ORDER = ["拳击", "帆船", "摔跤-自由式", "柔道", "跆拳道", "场地自行车", "射箭",
               "皮划艇激流回旋", "柔术", "壁球", "网球", "电子竞技", "足球",
               "跆拳道-虚拟", "沙滩排球", "曲棍球"]
NOC_ORDER = ["中国", "日本", "韩国", "印度", "乌兹别克斯坦", "哈萨克斯坦", "伊朗",
             "泰国", "马来西亚", "中国台北", "阿联酋", "菲律宾", "沙特阿拉伯"]

THIN = Border(*[Side(style="thin", color="999999")] * 4)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)


def style_header(ws, row, cols):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN


def main():
    wb = Workbook()
    gen_time = "2026-10-07"

    # ============ Sheet1: 明细 ============
    ws = wb.active
    ws.title = "明细"
    ws["A1"] = "第20届亚运会（爱知·名古屋 2026）— 2026-10-02 金牌明细（第13比赛日，全天已结束）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws.append([])
    ws.append(["序号", "大项", "分项", "冠军代表团", "冠军/队伍", "备注"])
    style_header(ws, 3, range(1, 7))
    for r in ROWS:
        ws.append(list(r))
    n = len(ROWS)
    for row in ws.iter_rows(min_row=4, max_row=3 + n, max_col=6):
        for cell in row:
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([6, 15, 17, 13, 30, 36], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    # ============ Sheet2: 透视 ============
    pv = wb.create_sheet("透视")
    pv["A1"] = "透视汇总 — 2026-10-02 金牌（所有汇总数字均为公式，引用『明细』页）"
    pv["A1"].font = TITLE_FONT
    pv.merge_cells("A1:D1")

    D = "明细!"
    rng_sport = f"{D}$B$4:$B${3+n}"
    rng_noc = f"{D}$D$4:$D${3+n}"

    pv["A3"] = "一、按大项汇总"
    pv["A3"].font = Font(bold=True, size=12)
    pv["A4"] = "大项"; pv["B4"] = "金牌数"
    style_header(pv, 4, (1, 2))
    r = 5
    for s in SPORT_ORDER:
        pv.cell(row=r, column=1, value=s)
        pv.cell(row=r, column=2, value=f'=COUNTIF({rng_sport},A{r})')
        r += 1
    last_sport_row = r - 1
    pv.cell(row=r, column=1, value="合计")
    pv.cell(row=r, column=1).font = Font(bold=True)
    pv.cell(row=r, column=2, value=f"=SUM(B5:B{last_sport_row})")
    pv.cell(row=r, column=2).font = Font(bold=True)
    sport_total_row = r

    r += 2
    pv.cell(row=r, column=1, value="二、按代表团（国家/地区）汇总").font = Font(bold=True, size=12)
    r += 1
    pv.cell(row=r, column=1, value="代表团")
    pv.cell(row=r, column=2, value="金牌数")
    style_header(pv, r, (1, 2))
    r += 1
    noc_first = r
    for noc in NOC_ORDER:
        pv.cell(row=r, column=1, value=noc)
        pv.cell(row=r, column=2, value=f'=COUNTIF({rng_noc},A{r})')
        r += 1
    noc_last = r - 1
    pv.cell(row=r, column=1, value="合计")
    pv.cell(row=r, column=1).font = Font(bold=True)
    pv.cell(row=r, column=2, value=f"=SUM(B{noc_first}:B{noc_last})")
    pv.cell(row=r, column=2).font = Font(bold=True)
    noc_total_row = r

    r += 2
    pv.cell(row=r, column=1, value="三、口径校验").font = Font(bold=True, size=12)
    r += 1
    pv.cell(row=r, column=1, value="校验项")
    pv.cell(row=r, column=2, value="数值")
    style_header(pv, r, (1, 2))
    r += 1
    pv.cell(row=r, column=1, value="明细行数（COUNTA 公式）")
    pv.cell(row=r, column=2, value=f"=COUNTA({D}$A$4:$A${3+n})")
    c0 = r
    r += 1
    pv.cell(row=r, column=1, value="按大项合计")
    pv.cell(row=r, column=2, value=f"=B{sport_total_row}")
    r += 1
    pv.cell(row=r, column=1, value="按代表团合计")
    pv.cell(row=r, column=2, value=f"=B{noc_total_row}")
    r += 1
    pv.cell(row=r, column=1, value="任务口径应产金数（用户给定）")
    pv.cell(row=r, column=2, value=51)
    r += 1
    pv.cell(row=r, column=1, value="四方是否一致")
    pv.cell(row=r, column=2,
             value=f'=IF(AND(B{c0}=B{c0+1},B{c0+1}=B{c0+2},B{c0+2}=B{c0+3}),"一致","不一致")')

    for rr in range(3, r + 1):
        for cc in (1, 2):
            pv.cell(row=rr, column=cc).border = THIN
    pv.column_dimensions["A"].width = 42
    pv.column_dimensions["B"].width = 14

    # ============ Sheet3: 口径说明 ============
    kj = wb.create_sheet("口径说明")
    kj["A1"] = "口径说明"
    kj["A1"].font = TITLE_FONT
    kj.column_dimensions["A"].width = 110
    lines = [
        "1. 统计对象：第20届亚洲运动会（爱知·名古屋 2026）2026-10-02（第13比赛日，星期五）当日决出的全部金牌；当日为本届产金最多的一天，按官方赛程全天共决出 51 枚金牌。",
        "2. 当日为全天已结束状态，按赛后最终成绩采集（多来源交叉核验，见『数据来源』页）。",
        "3. 金牌归属判定：以各单项决赛最终结果为准；团体/双人项目按代表团计 1 枚金牌。",
        "4. 代表团按 NOC 口径统计；中国香港、中国台北等均以代表团名称列示（中国香港当日未夺金）。",
        "5. 当日金牌按大项分布（与官方当日赛程一致）：拳击11、帆船5、摔跤5、柔道5、跆拳道4、场地自行车4、射箭3、皮划艇激流回旋2、柔术3、壁球2、网球2、电子竞技1、足球1、虚拟跆拳道1（本届新增小项）、沙滩排球1、曲棍球1，合计 51。",
        "6. 帆船 iQFoil（水翼帆板）男、女两级：因10月3日风力不足取消后续比赛，按10月2日末轮结束后名次定格判金（奖牌仪式于次日举行），故计入10月2日产出；其余帆船三级（ILCA7、49erFX、470混合）当日决赛正常进行。",
        "7. 当日部分代表团产金数（供核对）：日本13、中国9（含上述帆船2项按名次定格判金；若按奖牌仪式当日口径中国为7）、韩国3、印度6、乌兹别克斯坦4、哈萨克斯坦3、伊朗3、马来西亚3、泰国2、阿联酋2、中国台北1、菲律宾1、沙特阿拉伯1。",
        "8. 透视页所有汇总数字均为 Excel 公式（COUNTIF/SUM/COUNTA/IF）引用『明细』页实时计算，严禁且未手写任何汇总常数；透视页校验区对『明细行数、按大项合计、按代表团合计、任务口径51』四方相互核对，结果应为『一致』。",
        "9. 任务口径应产生金牌总数 = 51（用户给定）。",
        f"10. 工作簿生成时间：{gen_time}（UTC）；生成环境：Python 3 + openpyxl。",
    ]
    for i, t in enumerate(lines, start=3):
        kj.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")

    # ============ Sheet4: 数据来源 ============
    src = wb.create_sheet("数据来源")
    src["A1"] = "数据来源（2026-10-07 采集，赛后多源交叉核验）"
    src["A1"].font = TITLE_FONT
    src.append([])
    src.append(["#", "来源", "说明", "链接"])
    style_header(src, 3, range(1, 5))
    sources = [
        (1, "名古屋亚运会官方成绩系统（Bornan）", "各单项决赛官方结果",
         "https://results.asiangames2026.org/"),
        (2, "爱知·名古屋 2026 组委会官网", "官方赛程：10月2日 16 个大项决出 51 金及其分布",
         "https://www.aichi-nagoya2026.org/sports/"),
        (3, "日本奥委会（JOC）当日结果页", "日本队 10月2日全部结果（当日 13 金逐项核实）",
         "https://www.joc.or.jp/games/asia/2026/schedule/result_1002.html"),
        (4, "World Boxing（世界拳击）", "拳击 11 个级别决赛全部结果",
         "https://worldboxing.org/eleven-champions-crowned-as-india-top-aichi-nagoya-2026-asian-games-boxing-medal-table/"),
        (5, "Olympics.com 亚运专题", "印度当日 6 金战报（摔跤/拳击/女曲/射箭女团）",
         "https://www.olympics.com/en/news/asian-games-2026-live-scores-updates-results-india-october-2"),
        (6, "KhelNow 第13日奖牌综述", "当日各国金牌归属（拳击/柔道/跆拳道/自行车/皮划艇/网球/电竞等）",
         "https://khelnow.com/asian-games/medal-tally-october-2-asian-games-202610"),
        (7, "央视网亚运专题 / 搜狐体育", "中国代表团当日夺金明细与奖牌榜日增量",
         "https://yayun.cctv.com/2026/dateschedule/index.shtml"),
        (8, "United World Wrestling（UWW）", "摔跤第3日（10月2日）5 个级别决赛结果",
         "https://uww.org/articles/asian-games-2026-wrestling-day-3-highlights-susaki-sujeet-zare-win-gold"),
        (9, "共同社（Kyodo）", "女足决赛（日本点球胜朝鲜，三连冠）",
         "https://english.kyodonews.net/articles/-/87069"),
        (10, "World Squash", "壁球男/女团体决赛（马来西亚双金）",
         "https://www.worldsquash.sport/historic-gold-medal-double-for-malaysia-at-2026-asian-games-teams/"),
        (11, "Asian Sailing / Wikipedia 帆船页", "帆船 5 个级别结果及 iQFoil 名次定格说明",
         "https://www.asiansailing.org/asian-games-final-day-report/"),
        (12, "Sharjah24 / Khaleej Times（阿联酋）", "柔术当日金牌（阿联酋 2 金）",
         "https://sharjah24.ae/en/Articles/2026/10/02/AL011"),
        (13, "SBS（韩国）/ QazInform（哈萨克斯坦）", "柔术男子-94公斤级决赛（沙特阿拉伯夺金）",
         "https://news.sbs.co.kr/english/article.do?news_id=N1008780926"),
        (14, "ESPN 菲律宾代表团追踪", "虚拟跆拳道混合个人金牌（菲律宾）",
         "https://www.espn.com/espn/story/_/id/49945208/asian-games-2026-aichi-nagoya-team-philippines-results-schedule-medal-tracker"),
        (15, "Wikipedia: 2026 Asian Games medal table", "各代表团累计金牌数日增量交叉核验",
         "https://en.wikipedia.org/wiki/2026_Asian_Games_medal_table"),
    ]
    for s in sources:
        src.append(list(s))
    for row in src.iter_rows(min_row=4, max_row=3 + len(sources), max_col=4):
        for cell in row:
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate([5, 30, 46, 58], start=1):
        src.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
